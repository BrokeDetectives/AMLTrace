"""
Visual + geometric QA for the deck, without LibreOffice.

Reads the generated .pptx back through python-pptx and re-draws every slide with
PIL at the recorded geometry. It is an approximation of PowerPoint's renderer,
but it is driven entirely by what is actually in the file, so it catches the
defects that matter: text overflowing its box, shapes off the slide, overlaps,
and margin violations.

    python deck/qa_render.py deck/AMLTrace_Pitch_Deck.pptx
"""
import math, os, sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

import matplotlib.font_manager as fm

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "AMLTrace_Pitch_Deck.pptx")
SCALE = 110                      # px per inch
MARGIN_MIN = 0.45                # inches


def font_path(bold=False):
    for name in (["Arial Bold", "DejaVu Sans Bold"] if bold else ["Arial", "DejaVu Sans"]):
        try:
            return fm.findfont(fm.FontProperties(family=name.replace(" Bold", ""),
                                                 weight="bold" if bold else "normal"),
                               fallback_to_default=True)
        except Exception:
            continue
    return fm.findfont(fm.FontProperties())


FP_REG, FP_BOLD = font_path(False), font_path(True)
_cache = {}


def pil_font(size_pt, bold):
    key = (round(size_pt, 1), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FP_BOLD if bold else FP_REG, max(6, int(size_pt * SCALE / 72)))
    return _cache[key]


def emu_in(v):
    return 0 if v is None else Emu(v).inches


def rgb(c, default=(200, 200, 200)):
    try:
        if c and c.type is not None and c.rgb is not None:
            return tuple(int(str(c.rgb)[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        pass
    return default


def wrap(draw, txt, font, maxpx):
    """Greedy wrap; returns list of lines."""
    words, lines, cur = txt.split(), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font) <= maxpx or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def main():
    prs = Presentation(DECK)
    SW, SH = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    issues = []
    out_paths = []

    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (int(SW * SCALE), int(SH * SCALE)), (7, 12, 24))
        d = ImageDraw.Draw(img, "RGBA")

        # slide background if set
        try:
            bg = slide.background.fill
            if bg.type is not None:
                d.rectangle([0, 0, img.width, img.height], fill=rgb(bg.fore_color, (7, 12, 24)))
        except Exception:
            pass

        boxes = []
        for sh in slide.shapes:
            x, y = emu_in(sh.left), emu_in(sh.top)
            w, h = emu_in(sh.width), emu_in(sh.height)
            px = [x * SCALE, y * SCALE, (x + w) * SCALE, (y + h) * SCALE]

            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                issues.append(f"slide {idx}: shape off-slide  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}")
            elif not sh.has_text_frame or sh.text_frame.text.strip():
                if x < MARGIN_MIN - 0.02 and w < SW - 1:
                    issues.append(f"slide {idx}: left margin {x:.2f}in < {MARGIN_MIN}in")
                if y < MARGIN_MIN - 0.02 and h < SH - 1:
                    issues.append(f"slide {idx}: top margin {y:.2f}in < {MARGIN_MIN}in")
                if SW - (x + w) < MARGIN_MIN - 0.02 and w < SW - 1:
                    issues.append(f"slide {idx}: right margin {SW-(x+w):.2f}in < {MARGIN_MIN}in")
                if SH - (y + h) < MARGIN_MIN - 0.02 and h < SH - 1:
                    issues.append(f"slide {idx}: bottom margin {SH-(y+h):.2f}in < {MARGIN_MIN}in")

            if sh.shape_type is not None and str(sh.shape_type).startswith("PICTURE"):
                try:
                    im = Image.open(io_bytes(sh)).convert("RGB")
                    im = im.resize((max(1, int(w * SCALE)), max(1, int(h * SCALE))))
                    img.paste(im, (int(x * SCALE), int(y * SCALE)))
                except Exception:
                    d.rectangle(px, outline=(90, 110, 150), width=2)
                boxes.append(("picture", px, ""))
                continue

            if sh.has_text_frame and not sh.text_frame.text.strip() or not sh.has_text_frame:
                # a drawn shape (card, circle, badge)
                try:
                    fill = rgb(sh.fill.fore_color, None) if sh.fill.type is not None else None
                except Exception:
                    fill = None
                try:
                    ln = rgb(sh.line.color, None) if sh.line.fill.type is not None else None
                except Exception:
                    ln = None
                name = str(sh.shape_type)
                if "OVAL" in name:
                    d.ellipse(px, fill=fill, outline=ln, width=2)
                else:
                    d.rounded_rectangle(px, radius=8, fill=fill, outline=ln, width=2)
                boxes.append(("shape", px, ""))
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue

            # text
            tf = sh.text_frame
            cy = y * SCALE + 2
            maxpx = w * SCALE - 4
            needed = 0.0
            for p in tf.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs:
                    cy += 6
                    needed += 6
                    continue
                size = max([(r.font.size.pt if r.font.size else 14) for r in runs])
                bold = any(bool(r.font.bold) for r in runs)
                col = None
                for r in runs:
                    col = rgb(r.font.color, None)
                    if col:
                        break
                f = pil_font(size, bold)
                line_h = size * SCALE / 72 * (float(p.line_spacing) if isinstance(p.line_spacing, float) else 1.2)
                txt = "".join(r.text for r in runs)
                lines = wrap(d, txt, f, maxpx)
                for ln_txt in lines:
                    tw = d.textlength(ln_txt, font=f)
                    tx = x * SCALE
                    try:
                        al = str(p.alignment)
                    except Exception:
                        al = "LEFT"
                    if "CENTER" in al:
                        tx = x * SCALE + (w * SCALE - tw) / 2
                    d.text((tx, cy), ln_txt, font=f, fill=col or (232, 238, 251))
                    cy += line_h
                    needed += line_h
                sa = p.space_after.pt if p.space_after else 0
                cy += sa * SCALE / 72
                needed += sa * SCALE / 72
            avail = h * SCALE
            if needed > avail + 6:
                issues.append(f"slide {idx}: TEXT OVERFLOW {(needed-avail)/SCALE:.2f}in - "
                              f"\"{tf.text.strip()[:58]}\"")
            boxes.append(("text", [x * SCALE, y * SCALE, (x + w) * SCALE, y * SCALE + needed],
                          tf.text.strip()[:40]))

        # text-on-text overlap check
        texts = [b for b in boxes if b[0] == "text"]
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i][1], texts[j][1]
                ox = min(a[2], b[2]) - max(a[0], b[0])
                oy = min(a[3], b[3]) - max(a[1], b[1])
                if ox > 10 and oy > 10:
                    issues.append(f"slide {idx}: text overlap "
                                  f"\"{texts[i][2]}\" x \"{texts[j][2]}\" "
                                  f"({ox/SCALE:.2f}in x {oy/SCALE:.2f}in)")

        p = os.path.join(HERE, f"qa-{idx:02d}.png")
        img.save(p)
        out_paths.append(p)

    print(f"Rendered {len(out_paths)} slides -> deck/qa-NN.png")
    if issues:
        print(f"\n{len(issues)} issue(s):")
        for i in issues:
            print("  -", i)
    else:
        print("\nNo geometry or overflow issues detected.")
    return issues


def io_bytes(shape):
    import io
    return io.BytesIO(shape.image.blob)


if __name__ == "__main__":
    main()
